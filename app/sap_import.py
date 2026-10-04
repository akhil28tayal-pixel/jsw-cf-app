"""
Import logic for the two SAP exports JSW gives you:

  Sale_*.xlsx        -> one row per billing line -> becomes a Billing entry
  Material_In_*.xlsx -> one row per truck GRN     -> becomes a GRN entry

Both files report quantities in METRIC TONNES, not bags (confirmed from the
sample files: a QTY of 1.25 MT divides evenly into 25 bags at 50kg/bag,
which a bags-based quantity never would). We convert MT -> bags using the
same bag_weight_mt setting the rest of the app uses, so everything stays in
bags internally.

Column names are matched by header text (case/space-insensitive) rather than
position, since SAP export column order can shift between extracts.
"""
import datetime as dt
import io
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from app import models, crud


def _norm(s) -> str:
    return str(s or "").strip().lower()


def _read_sheet_as_dicts(file_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_norm(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(header, r))
        out.append(d)
    return out


def _to_date(value) -> Optional[dt.date]:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return dt.datetime.strptime(str(value).strip(), "%d.%m.%Y").date()
    except ValueError:
        return None


def _to_float(value) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


@dataclass
class ImportResult:
    rows_total: int = 0
    rows_imported: int = 0
    rows_skipped_duplicate: int = 0
    rows_skipped_unmapped: int = 0
    rows_skipped_pending: int = 0
    rows_skipped_other: int = 0
    unmapped_codes: set = field(default_factory=set)
    messages: list = field(default_factory=list)
    # (table name, row id) for every row this run created, so log_import can
    # stamp them with the log id and the import becomes undoable.
    created_rows: list = field(default_factory=list)

    def summary(self) -> str:
        parts = [f"{self.rows_imported} imported"]
        if self.rows_skipped_duplicate:
            parts.append(f"{self.rows_skipped_duplicate} already imported")
        if self.rows_skipped_unmapped:
            parts.append(f"{self.rows_skipped_unmapped} unmapped product code")
        if self.rows_skipped_pending:
            parts.append(f"{self.rows_skipped_pending} still in transit")
        if self.rows_skipped_other:
            parts.append(f"{self.rows_skipped_other} skipped (see details)")
        return ", ".join(parts) + f" out of {self.rows_total} rows."

    def details_text(self) -> str:
        lines = list(self.messages)
        if self.unmapped_codes:
            lines.append("Unmapped SAP codes seen: " + ", ".join(sorted(self.unmapped_codes)))
        return "\n".join(lines)


def resolve_product(db: Session, sap_code: str, sap_description: str = "") -> Optional[models.Product]:
    """Look up (or auto-register as unmapped) a SAP material/brand code."""
    if not sap_code:
        return None
    row = db.query(models.SapProductMap).filter(models.SapProductMap.sap_code == sap_code).first()
    if row is None:
        # First time we've ever seen this code — record it as unmapped so it
        # shows up on the mapping page for an admin to assign, instead of
        # silently failing every time this code appears.
        row = models.SapProductMap(sap_code=sap_code, sap_description=sap_description, product_id=None)
        db.add(row)
        db.commit()
        return None
    return row.product  # may be None if an admin explicitly marked it "ignore"


def get_or_create_dealer(db: Session, sap_code: str, name: str) -> models.Dealer:
    name = (name or "").strip() or f"Unknown Dealer ({sap_code})"
    if sap_code:
        d = db.query(models.Dealer).filter(models.Dealer.sap_code == sap_code).first()
        if d:
            return d
    d = db.query(models.Dealer).filter(models.Dealer.name == name).first()
    if d:
        if sap_code and not d.sap_code:
            d.sap_code = sap_code
            db.commit()
        return d
    d = models.Dealer(name=name, sap_code=sap_code or None)
    db.add(d)
    db.commit()
    db.refresh(d)
    return d


def get_or_create_transporter(db: Session, sap_code: str, name: str) -> Optional[models.Transporter]:
    name = (name or "").strip()
    if not name:
        return None
    if sap_code:
        t = db.query(models.Transporter).filter(models.Transporter.sap_code == sap_code).first()
        if t:
            return t
    t = db.query(models.Transporter).filter(models.Transporter.name == name).first()
    if t:
        if sap_code and not t.sap_code:
            t.sap_code = sap_code
            db.commit()
        return t
    t = models.Transporter(name=name, sap_code=sap_code or None)
    db.add(t)
    db.commit()
    db.refresh(t)
    return t


# ---------------------------------------------------------------------------
# Sale export -> Billing
# ---------------------------------------------------------------------------
def import_sale_file(db: Session, file_bytes: bytes, godown_id: int, user_id: Optional[int] = None) -> ImportResult:
    rows = _read_sheet_as_dicts(file_bytes)
    result = ImportResult()

    for row in rows:
        invoice_no = str(row.get("invoice no") or "").strip()
        if not invoice_no:
            continue  # blank trailer row
        result.rows_total += 1

        if str(row.get("cancelled invoice") or "").strip():
            result.rows_skipped_other += 1
            result.messages.append(f"Invoice {invoice_no}: cancelled in SAP, skipped.")
            continue

        existing = db.query(models.Billing).filter(models.Billing.invoice_no == invoice_no).first()
        if existing:
            result.rows_skipped_duplicate += 1
            continue

        sap_code = str(row.get("product") or "").strip()
        product = resolve_product(db, sap_code, str(row.get("product description") or ""))
        if product is None:
            result.rows_skipped_unmapped += 1
            result.unmapped_codes.add(sap_code)
            continue

        date = _to_date(row.get("invoice date"))
        if date is None:
            result.rows_skipped_other += 1
            result.messages.append(f"Invoice {invoice_no}: no invoice date, skipped.")
            continue

        qty_mt = _to_float(row.get("qty"))
        bw = crud.product_bag_weight_mt(db, product)  # 50 kg cement, 20 kg Microfine
        bags = round(qty_mt / bw) if bw else 0
        amount = _to_float(row.get("total value"))

        dealer_sap = str(row.get("sold to party") or "").strip()
        dealer_name = str(row.get("sold to name") or "").strip()
        dealer = get_or_create_dealer(db, dealer_sap, dealer_name)

        remarks_bits = []
        odn = row.get("odn no.")
        truck = row.get("truck no")
        if odn:
            remarks_bits.append(f"ODN {odn}")
        if truck:
            remarks_bits.append(f"Truck {truck}")
        remarks = ("SAP import: " + ", ".join(remarks_bits)) if remarks_bits else "SAP import"

        entry = models.Billing(
            date=date, godown_id=godown_id, invoice_no=invoice_no, dealer_id=dealer.id, product_id=product.id,
            bags=bags, amount=amount, remarks=remarks, created_by=user_id,
        )
        db.add(entry)
        try:
            db.commit()
            result.rows_imported += 1
            result.created_rows.append(("billing", entry.id))
        except IntegrityError:
            db.rollback()
            result.rows_skipped_duplicate += 1

    return result


# ---------------------------------------------------------------------------
# Material In export -> GRN
# ---------------------------------------------------------------------------
def import_material_in_file(db: Session, file_bytes: bytes, godown_id: int, user_id: Optional[int] = None) -> ImportResult:
    rows = _read_sheet_as_dicts(file_bytes)
    result = ImportResult()

    for row in rows:
        delivery = str(row.get("delivery") or "").strip()
        if not delivery:
            continue  # blank trailer row
        result.rows_total += 1

        status = str(row.get("status") or "").strip().lower()
        received_qty_mt = _to_float(row.get("received qty"))
        material_doc = str(row.get("material document") or "").strip()

        # Register the product code regardless of whether this row is ready
        # to import yet, so it can be mapped ahead of time — but only count
        # it as an "unmapped" skip once it's actually eligible for import.
        sap_code = str(row.get("material") or "").strip()
        product = resolve_product(db, sap_code, str(row.get("description") or ""))

        if status != "inward" or received_qty_mt <= 0 or not material_doc:
            result.rows_skipped_pending += 1
            continue

        existing = db.query(models.GRN).filter(models.GRN.sap_grn_no == material_doc).first()
        if existing:
            result.rows_skipped_duplicate += 1
            continue

        if product is None:
            result.rows_skipped_unmapped += 1
            result.unmapped_codes.add(sap_code)
            continue

        date = _to_date(row.get("posting date")) or _to_date(row.get("act. gds mvmnt date"))
        if date is None:
            result.rows_skipped_other += 1
            result.messages.append(f"Material Doc {material_doc}: no usable date, skipped.")
            continue

        invoice_qty_mt = _to_float(row.get("actual delivery qty"))
        bw = crud.product_bag_weight_mt(db, product)  # 50 kg cement, 20 kg Microfine
        bags_invoice = round(invoice_qty_mt / bw) if bw else 0
        bags_received = round(received_qty_mt / bw) if bw else 0

        remarks_bits = []
        lr = row.get("lr number")
        transporter = row.get("transporter")
        pending = _to_float(row.get("pending qty"))
        if transporter:
            remarks_bits.append(f"Transporter {transporter}")
        if lr:
            remarks_bits.append(f"LR {lr}")
        if pending:
            remarks_bits.append(f"{pending:.2f} MT still pending on this delivery")
        remarks = ("SAP import: " + ", ".join(remarks_bits)) if remarks_bits else "SAP import"

        entry = models.GRN(
            date=date, godown_id=godown_id, sap_grn_no=material_doc,
            invoice_no=str(row.get("odn no") or "").strip() or None,
            vehicle_no=str(row.get("truck number") or "").strip() or None,
            product_id=product.id, bags_invoice=bags_invoice, bags_received=bags_received,
            source_plant=str(row.get("supplying plant") or "").strip() or None,
            remarks=remarks, created_by=user_id,
        )
        db.add(entry)
        try:
            db.commit()
            result.rows_imported += 1
            result.created_rows.append(("grn", entry.id))
        except IntegrityError:
            db.rollback()
            result.rows_skipped_duplicate += 1

    return result


def log_import(db: Session, import_type: str, filename: str, user_id: Optional[int], result: ImportResult,
                godown_id: Optional[int] = None):
    log = models.ImportLog(
        import_type=import_type, godown_id=godown_id, filename=filename, imported_by=user_id,
        rows_total=result.rows_total, rows_imported=result.rows_imported,
        rows_skipped_duplicate=result.rows_skipped_duplicate,
        rows_skipped_unmapped=result.rows_skipped_unmapped,
        rows_skipped_pending=result.rows_skipped_pending,
        rows_skipped_other=result.rows_skipped_other,
        details=result.details_text(),
    )
    db.add(log)
    db.commit()

    # Now that the log has an id, point every row this run created at it. The
    # rows have to be created first (the importer commits each one to catch
    # duplicates), so the link can only be made here.
    for table, row_id in result.created_rows:
        model = {"billing": models.Billing, "grn": models.GRN,
                 "sap_stock_snapshot": models.SapStockSnapshot}.get(table)
        if model is None:
            continue
        db.query(model).filter(model.id == row_id).update({"import_log_id": log.id})
    if result.created_rows:
        db.commit()
    return log
