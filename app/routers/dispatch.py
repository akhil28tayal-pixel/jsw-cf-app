import datetime as dt
import io
import re

from fastapi import APIRouter, Request, Depends, Form, Query
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database import get_db
from app.auth import require_login
from app.flash import flash, get_flashed_messages
from app.godown_context import get_active_godown
from app import crud, models
from app.query_utils import parse_date, parse_int

from app.templating import templates

router = APIRouter()


def _filtered_dispatches(db: Session, godown_id, date_from, date_to, product_id, dealer_id,
                         transporter_id, vehicle_no):
    """The one filter definition, shared by the Dispatch page and its Excel
    export, so a downloaded file always contains exactly the rows the screen
    was showing (the page caps the display at 500; the export does not)."""
    q = db.query(models.Dispatch)
    if godown_id:
        q = q.filter(models.Dispatch.godown_id == godown_id)
    if date_from:
        q = q.filter(models.Dispatch.date >= date_from)
    if date_to:
        q = q.filter(models.Dispatch.date <= date_to)
    if product_id:
        q = q.filter(models.Dispatch.product_id == product_id)
    if dealer_id:
        q = q.filter(models.Dispatch.dealer_id == dealer_id)
    if transporter_id:
        q = q.filter(models.Dispatch.transporter_id == transporter_id)
    if vehicle_no:
        q = q.filter(models.Dispatch.vehicle_no.ilike(f"%{vehicle_no}%"))
    return q


def _districts_for(db: Session, godown_id):
    q = db.query(models.FreightRateCard.district).distinct()
    if godown_id:
        q = q.filter(models.FreightRateCard.godown_id == godown_id)
    return sorted({d for (d,) in q.all() if d})


@router.get("/dispatch")
def dispatch_page(request: Request, date_from: str = Query(None), date_to: str = Query(None),
                   product_id: str = Query(None), dealer_id: str = Query(None),
                   transporter_id: str = Query(None), vehicle_no: str = Query(None),
                   db: Session = Depends(get_db), user=Depends(require_login)):
    active_godown = get_active_godown(request, db)
    date_from = parse_date(date_from)
    date_to = parse_date(date_to)
    product_id = parse_int(product_id)
    dealer_id = parse_int(dealer_id)
    transporter_id = parse_int(transporter_id)
    q = _filtered_dispatches(db, active_godown.id if active_godown else None, date_from, date_to,
                             product_id, dealer_id, transporter_id, vehicle_no)
    entries = q.order_by(models.Dispatch.date.desc(), models.Dispatch.id.desc()).limit(500).all()
    districts = _districts_for(db, active_godown.id if active_godown else None)
    return templates.TemplateResponse(request, "dispatch.html", {
        "user": user, "flashes": get_flashed_messages(request),
        "products": crud.list_products(db), "dealers": crud.list_dealers(db),
        "transporters": crud.list_transporters(db), "entries": entries, "districts": districts,
        "date_from": date_from, "date_to": date_to, "sel_product": product_id,
        "sel_dealer": dealer_id, "sel_transporter": transporter_id, "vehicle_no": vehicle_no or "",
    })


@router.post("/dispatch/add")
def add_dispatch(request: Request,
                  date: dt.date = Form(...), dc_no: str = Form(""),
                  dealer_name: str = Form(""), new_dealer: str = Form(""),
                  destination: str = Form(""), district: str = Form(""), pincode: str = Form(""),
                  vehicle_no: str = Form(""), transporter_name: str = Form(""), new_transporter: str = Form(""),
                  product_id: int = Form(...), bags: float = Form(...), remarks: str = Form(""),
                  db: Session = Depends(get_db), user=Depends(require_login)):
    active_godown = get_active_godown(request, db)
    dealer_final_name = new_dealer.strip() or dealer_name.strip()
    if not dealer_final_name:
        flash(request, "Select a dealer or type a new dealer name.", "error")
        return RedirectResponse("/dispatch", status_code=303)
    dealer = crud.add_dealer(db, dealer_final_name)
    transporter = None
    transporter_final_name = new_transporter.strip() or transporter_name.strip()
    if transporter_final_name:
        transporter = crud.add_transporter(db, transporter_final_name)

    entry = models.Dispatch(
        date=date, godown_id=active_godown.id, dc_no=dc_no.strip() or None, dealer_id=dealer.id,
        destination=destination.strip() or None, district=district.strip() or None,
        pincode=pincode.strip() or None, vehicle_no=vehicle_no.strip() or None,
        transporter_id=transporter.id if transporter else None,
        product_id=product_id, bags=bags, remarks=remarks.strip() or None,
        created_by=user.id,
    )
    db.add(entry)
    db.commit()
    flash(request, f"Dispatch of {bags:.0f} bags to {dealer_final_name} saved for {active_godown.name}.")
    return RedirectResponse("/dispatch", status_code=303)


@router.get("/dispatch/{dispatch_id}/edit")
def edit_dispatch_form(dispatch_id: int, request: Request, db: Session = Depends(get_db),
                        user=Depends(require_login)):
    entry = db.get(models.Dispatch, dispatch_id)
    if not entry:
        flash(request, "Dispatch entry not found.", "error")
        return RedirectResponse("/dispatch", status_code=303)
    districts = _districts_for(db, entry.godown_id)
    return templates.TemplateResponse(request, "edit_dispatch.html", {
        "user": user, "flashes": get_flashed_messages(request), "entry": entry,
        "products": crud.list_products(db), "dealers": crud.list_dealers(db),
        "transporters": crud.list_transporters(db), "districts": districts,
    })


@router.post("/dispatch/{dispatch_id}/edit")
def edit_dispatch_submit(dispatch_id: int, request: Request,
                          date: dt.date = Form(...), dc_no: str = Form(""),
                          dealer_name: str = Form(""), new_dealer: str = Form(""),
                          destination: str = Form(""), district: str = Form(""), pincode: str = Form(""),
                          vehicle_no: str = Form(""), transporter_name: str = Form(""),
                          new_transporter: str = Form(""),
                          product_id: int = Form(...), bags: float = Form(...), remarks: str = Form(""),
                          db: Session = Depends(get_db), user=Depends(require_login)):
    entry = db.get(models.Dispatch, dispatch_id)
    if not entry:
        flash(request, "Dispatch entry not found.", "error")
        return RedirectResponse("/dispatch", status_code=303)

    dealer_final_name = new_dealer.strip() or dealer_name.strip()
    if not dealer_final_name:
        flash(request, "Select a dealer or type a new dealer name.", "error")
        return RedirectResponse(f"/dispatch/{dispatch_id}/edit", status_code=303)
    dealer = crud.add_dealer(db, dealer_final_name)
    transporter = None
    transporter_final_name = new_transporter.strip() or transporter_name.strip()
    if transporter_final_name:
        transporter = crud.add_transporter(db, transporter_final_name)

    # Note: godown is intentionally NOT editable here — a dispatch belongs to
    # whichever godown it physically left from. If it was logged under the
    # wrong godown, delete it and re-add it under the correct one instead.
    entry.date = date
    entry.dc_no = dc_no.strip() or None
    entry.dealer_id = dealer.id
    entry.destination = destination.strip() or None
    entry.district = district.strip() or None
    entry.pincode = pincode.strip() or None
    entry.vehicle_no = vehicle_no.strip() or None
    entry.transporter_id = transporter.id if transporter else None
    entry.product_id = product_id
    entry.bags = bags
    entry.remarks = remarks.strip() or None
    db.commit()
    flash(request, f"Dispatch #{dispatch_id} updated.")
    return RedirectResponse("/dispatch", status_code=303)


@router.post("/dispatch/{dispatch_id}/delete")
def delete_dispatch(dispatch_id: int, request: Request, db: Session = Depends(get_db),
                     user=Depends(require_login)):
    entry = db.get(models.Dispatch, dispatch_id)
    if not entry:
        flash(request, "Dispatch entry not found.", "error")
        return RedirectResponse("/dispatch", status_code=303)
    summary = f"{entry.bags:.0f} bags of {entry.product.name} to {entry.dealer.name} on {entry.date}"
    db.delete(entry)
    db.commit()
    flash(request, f"Deleted dispatch: {summary}. Stock and the dealer advance/hold report have updated automatically.")
    return RedirectResponse("/dispatch", status_code=303)



EXPORT_COLUMNS = [
    ("Date", 12), ("DC/Invoice No.", 16), ("Dealer", 28), ("Destination", 22), ("District", 16),
    ("Pincode", 10), ("Vehicle No.", 14), ("Transporter", 24), ("Product", 18),
    ("Bags", 10), ("MT", 10), ("Remarks", 30),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-") or "all"


@router.get("/dispatch/export")
def export_dispatch_excel(request: Request, date_from: str = Query(None), date_to: str = Query(None),
                          product_id: str = Query(None), dealer_id: str = Query(None),
                          transporter_id: str = Query(None), vehicle_no: str = Query(None),
                          db: Session = Depends(get_db), user=Depends(require_login)):
    """Download the currently filtered dispatch list as .xlsx.

    The Filter form posts here through `formaction`, so whatever is on screen
    is what lands in the file — no second set of filter controls to keep in
    sync. Unlike the page, the export is not capped at 500 rows.
    """
    active_godown = get_active_godown(request, db)
    date_from = parse_date(date_from)
    date_to = parse_date(date_to)
    product_id = parse_int(product_id)
    dealer_id = parse_int(dealer_id)
    transporter_id = parse_int(transporter_id)
    rows = _filtered_dispatches(
        db, active_godown.id if active_godown else None, date_from, date_to,
        product_id, dealer_id, transporter_id, vehicle_no,
    ).order_by(models.Dispatch.date, models.Dispatch.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Dispatch"
    ws.append([label for label, _ in EXPORT_COLUMNS])
    for idx, (label, width) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width

    total_bags = 0.0
    total_mt = 0.0
    for d in rows:
        # Weight has to use the bag size of the product actually on the truck —
        # 400 bags of Microfine is 8 MT, not 20 MT.
        mt = (d.bags or 0) * crud.product_bag_weight_mt(db, d.product)
        total_bags += d.bags or 0
        total_mt += mt
        ws.append([
            d.date, d.dc_no or "", d.dealer.name if d.dealer else "", d.destination or "",
            d.district or "", d.pincode or "", d.vehicle_no or "",
            d.transporter.name if d.transporter else "", d.product.name if d.product else "",
            d.bags or 0, round(mt, 3), d.remarks or "",
        ])

    last_row = ws.max_row
    if rows:
        total_row = last_row + 1
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=10, value=total_bags).font = Font(bold=True)
        ws.cell(row=total_row, column=11, value=round(total_mt, 3)).font = Font(bold=True)
        for col in range(1, len(EXPORT_COLUMNS) + 1):
            ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor="DDEBF7")

    for row in ws.iter_rows(min_row=2, max_row=last_row):
        row[0].number_format = "dd-mmm-yy"
        row[9].number_format = "#,##0"
        row[10].number_format = "#,##0.000"
    if last_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{last_row}"
    ws.freeze_panes = "A2"

    parts = ["dispatch", _slug(active_godown.name if active_godown else "")]
    if date_from or date_to:
        parts.append(f"{date_from or 'start'}_to_{date_to or 'today'}")
    else:
        parts.append(dt.date.today().isoformat())
    filename = "_".join(parts) + ".xlsx"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
